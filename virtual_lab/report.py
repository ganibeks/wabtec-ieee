"""Downloadable test report for one module: HTML (print to PDF from the browser), CSV and Excel.

Every value comes from the tester tiles, which come from simulation result files; nothing is computed here.
"""
from __future__ import annotations

import csv
import html
import io
from datetime import datetime

from virtual_lab import datasets as D
from virtual_lab.tester import Tile

STATE_TXT = {"IN_LIMIT": "IN LIMIT", "OUT_OF_LIMIT": "OUT OF LIMIT", "NO_LIMIT": "NO APPROVED LIMIT",
             "NOT_TESTED": "NOT TESTED"}
COLUMNS = ["Check", "Value", "Unit", "Limit (source)", "Result", "In verdict", "How obtained", "Provenance"]


def _val(t: Tile) -> str:
    if t.value is None:
        return "—"
    return f"{t.value:.6g}" if t.state == "OUT_OF_LIMIT" else f"{t.value:.4g}"


def rows(tiles: list[Tile], excluded: list[Tile]) -> list[list[str]]:
    out = [[t.label, _val(t), t.unit, t.limit, STATE_TXT.get(t.state, t.state), "yes", t.note, t.provenance] for t in tiles]
    out += [[t.label, _val(t), t.unit, t.limit, STATE_TXT.get(t.state, t.state), "no (excluded)", t.note, t.provenance]
            for t in excluded]
    return out


def header(module_id: str, ds: D.Dataset, verdict: dict, when: datetime) -> list[tuple[str, str]]:
    return [("Module serial", module_id), ("Module under test", ds.label), ("Verdict", verdict["verdict"]),
            ("Verdict detail", verdict["detail"]), ("Hardware decision", verdict["hardware_decision"]),
            ("Data source", "SIMULATED — MATLAB R2026b / Simscape Electrical"), ("Simulation run", ds.folder.name),
            ("Report generated", when.strftime("%Y-%m-%d %H:%M:%S"))]


def to_csv(module_id, ds, verdict, tiles, excluded, when) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for k, v in header(module_id, ds, verdict, when):
        w.writerow([k, v])
    w.writerow([])
    w.writerow(COLUMNS)
    w.writerows(rows(tiles, excluded))
    return buf.getvalue().encode("utf-8-sig")      # BOM so Excel opens µ, Ω, ± correctly


def to_xlsx(module_id, ds, verdict, tiles, excluded, when) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Test report"
    red, green, amber = "D0021B", "1F7A45", "B06A00"
    ws.append(["IGBT Bench Tester — test report"]); ws["A1"].font = Font(bold=True, size=14, color=red)
    for k, v in header(module_id, ds, verdict, when):
        ws.append([k, v]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    vc = ws.cell(4, 2); vc.font = Font(bold=True, size=13, color=green if verdict["verdict"] == "PASS" else red)
    ws.append([])
    ws.append(COLUMNS)
    hr = ws.max_row
    for c in ws[hr]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=red)
    for r in rows(tiles, excluded):
        ws.append(r)
        res = ws.cell(ws.max_row, 5)
        res.font = Font(bold=True, color=green if r[4] == "IN LIMIT" else red if r[4] == "OUT OF LIMIT" else amber)
    for col, wdt in zip("ABCDEFGH", (26, 14, 8, 40, 18, 14, 80, 22)):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=hr + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_html(module_id, ds, verdict, tiles, excluded, when, safety=None) -> bytes:
    e = html.escape
    saf = ""
    if safety is not None:
        allok = all(ok for _, ok in safety)
        items = "".join(f"<tr><td>{e(lab)}</td><td><span class='tag {'ok' if ok else 'bad'}'>{'CONFIRMED' if ok else 'NOT CONFIRMED'}</span></td></tr>" for lab, ok in safety)
        saf = (f"<h2>Safety interlocks — operator confirmation</h2><table><tbody>{items}</tbody></table>"
               f"<p class='note'>{'All interlocks confirmed.' if allok else 'NOT ALL INTERLOCKS CONFIRMED — this report must not be used as a test record.'} "
               "Confirmed by the operator, not measured; the hard-wired interlocks act independently of the software.</p>")
    vcls = "pass" if verdict["verdict"] == "PASS" else "nopass" if verdict["reason"] == "FAILED" else "unv"
    head = "".join(f"<tr><th>{e(k)}</th><td>{e(v)}</td></tr>" for k, v in header(module_id, ds, verdict, when)
                   if k not in ("Verdict", "Verdict detail"))
    def trs(ts):
        out = ""
        for t in ts:
            st = {"IN_LIMIT": "ok", "OUT_OF_LIMIT": "bad"}.get(t.state, "na")
            out += (f"<tr><td><b>{e(t.label)}</b></td><td class='num'>{e(_val(t))} {e(t.unit)}</td><td>{e(t.limit)}</td>"
                    f"<td><span class='tag {st}'>{e(STATE_TXT.get(t.state, t.state))}</span></td>"
                    f"<td class='note'>{e(t.note)}<br><span class='prov'>{e(t.provenance)}</span></td></tr>")
        return out
    excl = ""
    if excluded:
        excl = (f"<h2>Reported but not in the verdict</h2><table><thead><tr><th>Check</th><th>Value</th><th>Limit</th>"
                f"<th>Result</th><th>Reason and source</th></tr></thead><tbody>{trs(excluded)}</tbody></table>")
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>Test report {e(module_id)}</title>
<style>
@page{{size:A4;margin:14mm 12mm}}
body{{font:13px/1.45 'Segoe UI',Roboto,Arial,sans-serif;color:#1a1414;margin:0;padding:24px;max-width:1000px}}
.bar{{height:5px;background:#d0021b;margin:-24px -24px 18px}}
.eb{{font:600 11px Arial,sans-serif;letter-spacing:.14em;text-transform:uppercase;color:#d0021b}}
h1{{font:800 28px/1.1 Arial,sans-serif;margin:4px 0 14px}} h1 em{{font-style:normal;color:#d0021b}}
h2{{font:700 15px Arial,sans-serif;margin:22px 0 8px;text-transform:uppercase;letter-spacing:.06em;color:#d0021b}}
.verdict{{display:flex;gap:22px;align-items:center;padding:14px 18px;border:1px solid #e8e1e1;margin-bottom:14px}}
.verdict b{{font:800 36px/1 Arial,sans-serif;letter-spacing:.03em}}
.verdict.pass{{border-left:8px solid #1f7a45}} .verdict.pass b{{color:#1f7a45}}
.verdict.nopass{{background:#d0021b;color:#fff;border-color:#d0021b}}
.verdict.unv{{border-left:8px solid #b06a00}} .verdict.unv b{{color:#b06a00}}
table{{border-collapse:collapse;width:100%;margin-bottom:6px}}
th,td{{text-align:left;vertical-align:top;padding:6px 8px;border-bottom:1px solid #e8e1e1}}
thead th{{font-size:11px;text-transform:uppercase;letter-spacing:.06em;color:#6b5f5f;border-bottom:2px solid #1a1414}}
.meta th{{width:170px;color:#6b5f5f;font-weight:600}}
td.num{{font-family:Consolas,'IBM Plex Mono',monospace;white-space:nowrap}}
td.note{{font-size:11.5px;color:#4a3f3f}} .prov{{color:#8a7e7e;font-size:10.5px;text-transform:uppercase;letter-spacing:.05em}}
.tag{{font:700 10.5px Arial,sans-serif;padding:3px 6px;white-space:nowrap}}
.tag.ok{{background:#e8f5ed;color:#1f7a45}} .tag.bad{{background:#d0021b;color:#fff}} .tag.na{{background:#f1eded;color:#6b5f5f}}
.sign{{display:flex;gap:40px;margin-top:28px}} .sign div{{flex:1;border-top:1px solid #1a1414;padding-top:4px;font-size:11px;color:#6b5f5f}}
.foot{{margin-top:18px;font-size:11px;color:#6b5f5f}}
@media print{{body{{padding:0}} .bar{{margin:0 0 14px}} tr{{break-inside:avoid}}}}
</style></head><body>
<div class="bar"></div>
<div class="eb">NU × Wabtec Hackathon 2026 · pre-installation test cell</div>
<h1>IGBT Bench <em>Tester</em> — test report</h1>
<div class="verdict {vcls}"><b>{e(verdict['verdict'])}</b><div>{e(verdict['detail'])}<br>
<small>Simulation only, not a hardware decision (hardware decision: {e(verdict['hardware_decision'])}).</small></div></div>
<table class="meta">{head}</table>
<h2>Checks in the verdict</h2>
<table><thead><tr><th>Check</th><th>Value</th><th>Limit</th><th>Result</th><th>How obtained and source</th></tr></thead>
<tbody>{trs(tiles)}</tbody></table>
{excl}
{saf}
<div class="sign"><div>Operator</div><div>Date</div><div>Signature</div></div>
<div class="foot">Limits: Wabtec brief slides 5 and 8. Device: representative Mitsubishi CM1500HC-66XB, not a confirmed Wabtec part.
Every value comes from MATLAB/Simscape result files ({e(ds.folder.name)} and the linked static, diode and DESAT runs).
Generated {when.strftime('%Y-%m-%d %H:%M:%S')} by IGBT Bench Tester. To save as PDF: open in a browser, Ctrl+P, “Save as PDF”.</div>
</body></html>"""
    return page.encode("utf-8")
